"""convert_document: offline conversion between common document, data and image formats.

The original file is never changed. Reading goes through check_read_file, writing through check_write_path.
Supported pairs are listed in CONVERSIONS (the tool description is built from it)."""
import html as html_lib
import io
import json
import os
import re
import uuid
import zipfile

from langchain_core.tools import tool

from security import permissions
from security.settings import get_settings
from tools.common import fail, guarded, human_size, limits, ok
from tools.documents.others.extract import html_to_text
from security import guard
from security.guard import row
from tools.documents.files import file_reader as fr

IMAGES = {".png", ".jpg", ".webp", ".bmp", ".gif", ".tif"}   # canonical names (.jpeg = .jpg, .tiff = .tif)

# source extension -> target extensions
CONVERSIONS = {
    ".docx": {".txt", ".md"},
    ".pdf": {".txt", ".md"},
    ".txt": {".docx", ".pdf"},
    ".md": {".docx", ".pdf", ".txt"},
    ".html": {".txt"},
    ".csv": {".json", ".xlsx"},
    ".json": {".csv", ".xlsx"},
    ".xlsx": {".csv", ".json"},
}
for _img in IMAGES:
    CONVERSIONS[_img] = (IMAGES - {_img}) | {".pdf"}


def supported_text() -> str:
    return "; ".join(f"{src} -> {', '.join(sorted(dst))}" for src, dst in sorted(CONVERSIONS.items()))


def _cfg() -> dict:
    return get_settings().get("convert", {})


# ---------------------------------------------------------------- markdown helpers

INLINE = re.compile(r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*|`[^`\n]+`)")


def md_blocks(text: str, markdown: bool = True) -> list:
    """Tiny Markdown reader: ('h1'..'h6' | 'p' | 'li' | 'ol' | 'code', text). Plain text keeps its line breaks."""
    blocks, para, code, in_code = [], [], [], False

    def flush():
        if para:
            blocks.append(("p", (" " if markdown else "\n").join(para)))
            para.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if markdown and stripped.startswith("```"):
            if in_code:
                blocks.append(("code", "\n".join(code)))
                code.clear()
                in_code = False
            else:
                flush()
                in_code = True
            continue
        if in_code:
            code.append(line)
            continue
        if not stripped:
            flush()
            continue
        if markdown:
            m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if m:
                flush()
                blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
                continue
            m = re.match(r"^[-*+]\s+(.*)$", stripped)
            if m:
                flush()
                blocks.append(("li", m.group(1)))
                continue
            m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
            if m:
                flush()
                blocks.append(("ol", m.group(1)))
                continue
        para.append(stripped if markdown else line.rstrip())
    flush()
    if in_code:
        blocks.append(("code", "\n".join(code)))
    return blocks


def _inline_html(text: str) -> str:
    out = []
    for part in INLINE.split(html_lib.escape(text, quote=False)):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append(f"<b>{part[2:-2]}</b>")
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append(f"<code>{part[1:-1]}</code>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(f"<i>{part[1:-1]}</i>")
        else:
            out.append(part)
    return "".join(out)


def blocks_to_html(blocks: list) -> str:
    parts, open_list = [], None
    for kind, text in blocks:
        want = {"li": "ul", "ol": "ol"}.get(kind)
        if open_list and open_list != want:
            parts.append(f"</{open_list}>")
            open_list = None
        if want and not open_list:
            parts.append(f"<{want}>")
            open_list = want
        if kind.startswith("h"):
            parts.append(f"<{kind}>{_inline_html(text)}</{kind}>")
        elif want:
            parts.append(f"<li>{_inline_html(text)}</li>")
        elif kind == "code":
            parts.append(f"<pre>{html_lib.escape(text, quote=False)}</pre>")
        else:
            parts.append("<p>" + "<br>".join(_inline_html(line) for line in text.split("\n")) + "</p>")
    if open_list:
        parts.append(f"</{open_list}>")
    return "".join(parts)


def _add_runs(paragraph, text: str):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


# ---------------------------------------------------------------- converters: each writes to `out` (a temp path)

def _text_of(path: str) -> str:
    text, _, _ = fr._read_text(path)
    return text


def _docx_to_text(src, out, ext_out):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    fr_check = zipfile.ZipFile(src)
    unpacked = sum(i.file_size for i in fr_check.infolist())
    fr_check.close()
    if unpacked > limits()["max_docx_unpacked_mb"] * 1024 * 1024:
        raise ValueError("The .docx expands to more than the allowed size.")
    markdown = ext_out == ".md"
    doc = Document(src)
    lines = []

    def run_text(par):
        if not markdown:
            return par.text
        out_parts = []
        for r in par.runs:
            t = r.text
            if not t.strip():
                out_parts.append(t)
            elif r.bold and r.italic:
                out_parts.append(f"***{t.strip()}***")
            elif r.bold:
                out_parts.append(f"**{t.strip()}**")
            elif r.italic:
                out_parts.append(f"*{t.strip()}*")
            else:
                out_parts.append(t)
        return "".join(out_parts) or par.text

    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            text = run_text(block).strip()
            if not text:
                lines.append("")
                continue
            style = (block.style.name if block.style is not None else "") or ""
            if markdown and style == "Title":
                lines.append(f"# {text}")
            elif markdown and re.match(r"Heading \d", style):
                lines.append("#" * min(int(style.split()[-1]), 6) + f" {text}")
            elif markdown and "List Bullet" in style:
                lines.append(f"- {text}")
            elif markdown and "List Number" in style:
                lines.append(f"1. {text}")
            else:
                lines.append(text)
        elif isinstance(block, Table):
            rows = [[c.text.strip().replace("\n", " ").replace("|", "\\|") for c in r.cells] for r in block.rows]
            if not rows:
                continue
            if markdown:
                lines.append("")
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + " --- |" * len(rows[0]))
                lines.extend("| " + " | ".join(r) + " |" for r in rows[1:])
                lines.append("")
            else:
                lines.extend(" | ".join(r) for r in rows)
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"
    _write_text(out, text)
    return [f"{len(text):,} characters written."]


def _pdf_to_text(src, out, ext_out):
    import pymupdf

    markdown = ext_out == ".md"
    ocr_cfg = get_settings().get("ocr", {})
    ocr_left = ocr_cfg.get("max_pages", 5) if ocr_cfg.get("enabled", False) else 0
    max_pages = int(_cfg().get("max_pdf_pages", 200))
    parts, notes, ocr_pages, empty, ocr_error = [], [], 0, 0, None
    with pymupdf.open(src) as doc:
        if doc.needs_pass and not doc.authenticate(""):
            raise ValueError("This PDF is password protected.")
        total = doc.page_count
        for i in range(min(total, max_pages)):
            page = doc[i]
            text = page.get_text("text").strip()
            if not text and page.get_images() and ocr_left > 0 and not ocr_error:
                try:
                    text = fr._ocr_page(page, ocr_cfg)
                    ocr_left -= 1
                    ocr_pages += 1
                except Exception as exc:  # noqa: BLE001
                    ocr_error = str(exc).splitlines()[0][:120]
            if not text:
                empty += 1
            parts.append((f"## Page {i + 1}\n\n" if markdown else f"--- Page {i + 1} ---\n") + text)
    body = ("\n\n" if markdown else "\n\n").join(parts)
    _write_text(out, body.strip() + "\n")
    if total > max_pages:
        notes.append(f"Only the first {max_pages} of {total} pages were converted (convert.max_pdf_pages).")
    if ocr_pages:
        notes.append(f"OCR was used on {ocr_pages} scanned page(s); text may contain errors.")
    if ocr_error:
        notes.append(f"Some scanned pages could not be read (OCR: {ocr_error}).")
    if empty:
        notes.append(f"{empty} page(s) had no text.")
    return notes or [f"{min(total, max_pages)} page(s) converted."]


def _text_to_docx(src, out, ext_out):
    from docx import Document

    markdown = os.path.splitext(src)[1].lower() == ".md"
    doc = Document()
    for kind, text in md_blocks(_text_of(src), markdown=markdown):
        if kind.startswith("h"):
            level = int(kind[1])
            heading = doc.add_heading(level=level)
            _add_runs(heading, text)
        elif kind == "li":
            _add_runs(doc.add_paragraph(style="List Bullet"), text)
        elif kind == "ol":
            _add_runs(doc.add_paragraph(style="List Number"), text)
        elif kind == "code":
            par = doc.add_paragraph()
            par.add_run(text).font.name = "Consolas"
        else:
            lines = text.split("\n")
            par = doc.add_paragraph()
            for n, line in enumerate(lines):
                _add_runs(par, line)
                if n < len(lines) - 1:
                    par.add_run().add_break()
    doc.save(out)
    return []


def _text_to_pdf(src, out, ext_out):
    # Formatted PDF (title, headings, nested/numbered lists, tables, links, page numbers): see pdf_render.py
    from tools.documents.pdf.pdf_render import render_pdf

    markdown = os.path.splitext(src)[1].lower() == ".md"
    stem = os.path.splitext(os.path.basename(src))[0]
    notes = render_pdf(_text_of(src), out, markdown=markdown, fallback_title=stem)
    return notes + ["Characters outside the built-in fonts (some non-Latin scripts, e.g. Devanagari) may not display."]


def _html_to_txt(src, out, ext_out):
    with open(src, "r", encoding="utf-8", errors="replace") as fh:
        text = html_to_text(fh.read())
    _write_text(out, text + "\n")
    return ["Scripts and styles were removed."]


def _write_text(out: str, text: str):
    with open(out, "w", encoding="utf-8", newline="") as fh:
        fh.write(text)


def _md_to_txt(src, out, ext_out):
    text = _text_of(src)
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\*\*([^*\n]+)\*\*|\*([^*\n]+)\*|`([^`\n]+)`", lambda m: m.group(1) or m.group(2) or m.group(3), text)
    _write_text(out, text)
    return ["Markdown symbols were removed."]


# --- tables (csv / json / xlsx) through pandas

def _sanitize(df):
    """Cells that start with = + - @ can run as formulas when the file is opened in Excel: prefix them with a quote."""
    import pandas as pd

    def fix(v):
        if isinstance(v, str) and v[:1] in ("=", "@", "+", "-"):
            try:
                float(v)
                return v
            except ValueError:
                return "'" + v
        return v

    changed = 0
    for col in df.columns:
        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            before = df[col]
            df[col] = before.map(fix)
            changed += int((before != df[col]).sum())
    return df, changed


def _load_table(src: str):
    import pandas as pd

    ext = os.path.splitext(src)[1].lower()
    infer = bool(_cfg().get("infer_types", False))
    kw = {} if infer else {"dtype": str, "keep_default_na": False}
    notes = []
    if ext == ".csv":
        df = pd.read_csv(src, sep=None, engine="python", encoding="utf-8-sig", **kw)
    elif ext == ".xlsx":
        zipfile_ok = zipfile.ZipFile(src)
        unpacked = sum(i.file_size for i in zipfile_ok.infolist())
        zipfile_ok.close()
        if unpacked > limits()["max_docx_unpacked_mb"] * 1024 * 1024:
            raise ValueError("The .xlsx expands to more than the allowed size.")
        book = pd.ExcelFile(src, engine="openpyxl")
        df = book.parse(book.sheet_names[0], **kw)
        if len(book.sheet_names) > 1:
            notes.append(f"Only the first sheet '{book.sheet_names[0]}' was converted ({len(book.sheet_names)} sheets in the file).")
    else:  # json
        with open(src, "r", encoding="utf-8-sig") as fh:
            data = json.load(fh)
        if isinstance(data, dict):
            if data and all(isinstance(v, list) for v in data.values()):
                df = pd.DataFrame(data)
            else:
                lists = [v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
                df = pd.json_normalize(lists[0] if len(lists) == 1 else [data])
        elif isinstance(data, list):
            if data and all(isinstance(r, dict) for r in data):
                df = pd.json_normalize(data)
            elif data and all(isinstance(r, list) for r in data):
                df = pd.DataFrame(data[1:], columns=[str(c) for c in data[0]])
            else:
                df = pd.DataFrame({"value": data})
        else:
            raise ValueError("The JSON must be a list of records (or an object that contains one).")
        for col in df.columns:  # nested lists/objects become JSON text so they fit in one csv/xlsx cell
            if df[col].map(lambda v: isinstance(v, (list, dict))).any():
                df[col] = df[col].map(lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v)
    return df, notes


def _table_convert(src, out, ext_out):
    df, notes = _load_table(src)
    if ext_out == ".csv":
        df, changed = _sanitize(df)
        df.to_csv(out, index=False, encoding="utf-8-sig")
    elif ext_out == ".xlsx":
        df, changed = _sanitize(df)
        df.to_excel(out, index=False, engine="openpyxl")
    else:
        changed = 0
        df.to_json(out, orient="records", force_ascii=False, indent=2)
    if changed:
        notes.append(f"{changed} cell(s) starting with = + - or @ got a leading ' so Excel cannot run them as formulas.")
    return notes + [f"{len(df):,} row(s), {len(df.columns)} column(s)."]


# --- images

def _image_convert(src, out, ext_out):
    from PIL import Image, ImageOps

    max_pixels = int(_cfg().get("max_image_pixels", 100_000_000))
    with Image.open(src) as img:
        if img.width * img.height > max_pixels:
            raise ValueError(f"Image is {img.width}x{img.height}, over the pixel limit (convert.max_image_pixels).")
        frame = ImageOps.exif_transpose(img)
        fmt = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP", ".bmp": "BMP", ".gif": "GIF",
               ".tif": "TIFF", ".tiff": "TIFF", ".pdf": "PDF"}[ext_out]
        notes = []
        if fmt in ("JPEG", "BMP", "PDF") and (frame.mode in ("RGBA", "LA", "P") or "transparency" in frame.info):
            rgba = frame.convert("RGBA")
            flat = Image.new("RGB", rgba.size, (255, 255, 255))
            flat.paste(rgba, mask=rgba.split()[-1])
            frame = flat
            notes.append("Transparency was replaced with a white background.")
        elif fmt in ("JPEG", "PDF") and frame.mode != "RGB":
            frame = frame.convert("RGB")
        kwargs = {"quality": 92} if fmt in ("JPEG", "WEBP") else {}
        if fmt == "PDF":
            kwargs["resolution"] = 150.0
        frame.save(out, format=fmt, **kwargs)
    notes.append("Metadata such as EXIF/GPS is not copied.")
    return notes


def _dispatch(src_ext: str, dst_ext: str):
    if src_ext in IMAGES:
        return _image_convert
    if src_ext == ".docx":
        return _docx_to_text
    if src_ext == ".pdf":
        return _pdf_to_text
    if src_ext in (".txt", ".md") and dst_ext == ".docx":
        return _text_to_docx
    if src_ext in (".txt", ".md") and dst_ext == ".pdf":
        return _text_to_pdf
    if src_ext == ".md" and dst_ext == ".txt":
        return _md_to_txt
    if src_ext in (".html", ".htm"):
        return _html_to_txt
    return _table_convert


# ---------------------------------------------------------------- destination logic

def _norm_ext(value: str) -> str:
    value = (value or "").strip().lower()
    if not value:
        return ""
    value = value if value.startswith(".") else "." + value
    return {".jpeg": ".jpg", ".tiff": ".tif", ".htm": ".html"}.get(value, value)


def _canon(ext: str) -> str:
    return {".jpeg": ".jpg", ".tiff": ".tif", ".htm": ".html"}.get(ext, ext)


def _pick_destination(real_src: str, target_format: str, destination: str):
    """Return (destination_path, target_ext, notes)."""
    src_ext = os.path.splitext(real_src)[1].lower()
    stem = os.path.splitext(os.path.basename(real_src))[0]
    fmt = _norm_ext(target_format)
    notes = []
    destination = (destination or "").strip().strip('"').strip("'")

    if destination:
        if os.path.isdir(destination) or destination.endswith(("\\", "/")):
            if not fmt:
                raise permissions.PermissionDenied("Give target_format (for example 'pdf') or a destination file name with an extension.")
            return os.path.join(destination, stem + fmt), fmt, notes
        dest_ext = _norm_ext(os.path.splitext(destination)[1])
        if not dest_ext:
            raise permissions.PermissionDenied("The destination must be an existing folder or a file name with an extension.")
        if fmt and fmt != dest_ext:
            raise permissions.PermissionDenied(f"target_format '{fmt}' does not match the destination extension '{dest_ext}'.")
        return destination, dest_ext, notes

    if not fmt:
        raise permissions.PermissionDenied("Give target_format, for example 'pdf', 'docx', 'md', 'csv' or 'png'.")
    default = os.path.join(os.path.dirname(real_src), stem + fmt)
    try:
        permissions.check_write_path(default, check_ext=True)
        return default, fmt, notes
    except permissions.PermissionDenied:
        out_dir = str(_cfg().get("output_dir", "")).strip()
        if not out_dir:
            raise
        notes.append(f"The source folder is not writable, so the result goes to {out_dir}.")
        return os.path.join(out_dir, stem + fmt), fmt, notes


def _write_via_temp(dest: str, ext_out: str, producer):
    tmp = os.path.join(os.path.dirname(dest), f".{uuid.uuid4().hex[:8]}.tmp{ext_out}")
    try:
        notes = producer(tmp)
        os.replace(tmp, dest)
        return notes
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


@tool
@guarded("convert_document")
def convert_document(source: str, target_format: str = "", destination: str = "", overwrite: bool = False) -> str:
    """Convert ONE file to another format; the original is never changed. `source` is an absolute path inside an approved read folder. `target_format` is the new extension (e.g. 'pdf'). `destination` is optional: a folder or a full file name inside an allowed write folder (default: next to the source). Supported (from -> to): docx -> txt, md; pdf -> txt, md; txt/md -> docx, pdf (md -> txt too); html -> txt; csv <-> json; csv/json <-> xlsx (xlsx -> csv, json); images png/jpg/webp/bmp/gif/tif -> any other image type or pdf. Not supported: docx -> pdf, pptx, old .doc. A new file is created immediately; replacing an existing file needs overwrite=true and the user's approval."""
    cfg = _cfg()
    if not cfg.get("enabled", True):
        return fail("convert_document", "Conversion is turned off (convert.enabled: false in config.yaml).")
    real = permissions.check_read_file(source)
    src_ext = _canon(os.path.splitext(real)[1].lower())
    if src_ext not in CONVERSIONS:
        return fail("convert_document", f"Cannot convert '{src_ext or '(no extension)'}' files. Supported conversions: {supported_text()}")
    size = os.path.getsize(real)
    if size > int(cfg.get("max_source_mb", 20)) * 1024 * 1024:
        return fail("convert_document", f"File is {human_size(size)}, over the convert.max_source_mb limit.")

    dest_raw, dst_ext, notes = _pick_destination(real, target_format, destination)
    dst_ext = _canon(dst_ext)
    if dst_ext not in CONVERSIONS[src_ext]:
        return fail("convert_document", f"Cannot convert {src_ext} to {dst_ext}. From {src_ext} you can create: "
                                        f"{', '.join(sorted(CONVERSIONS[src_ext]))}. All supported: {supported_text()}")
    final = permissions.check_write_path(dest_raw, check_ext=True)
    if os.path.normcase(final) == os.path.normcase(real):
        return fail("convert_document", "The destination is the source file itself.")
    exists = os.path.lexists(final)
    if exists and os.path.isdir(final):
        return fail("convert_document", f"A folder with that name already exists: {final}")
    if exists and not overwrite:
        return fail("convert_document", f"Already exists: {final}. Pass overwrite=true to replace it (the user must approve), "
                                        "or choose another destination.")

    def action():
        src = permissions.check_read_file(source)
        out = permissions.check_write_path(dest_raw, check_ext=True)
        if os.path.lexists(out) and not overwrite:
            return fail("convert_document", f"Already exists: {out}")
        parent = os.path.dirname(out)
        if not os.path.isdir(parent):
            os.makedirs(parent, exist_ok=True)
        converter = _dispatch(src_ext, dst_ext)
        try:
            result_notes = _write_via_temp(out, dst_ext, lambda tmp: converter(src, tmp, dst_ext))
        except ImportError as exc:
            return fail("convert_document", f"Missing package ({exc.name}). Run: python -m pip install -r requirements.txt")
        except (ValueError, zipfile.BadZipFile, json.JSONDecodeError, UnicodeDecodeError) as exc:
            return fail("convert_document", f"Could not convert this file: {exc}")
        except OSError:
            raise
        except Exception as exc:  # noqa: BLE001
            return fail("convert_document", f"Conversion failed ({type(exc).__name__}: {exc}).")
        detail = f"{src_ext} -> {dst_ext}, {human_size(os.path.getsize(out))} (original: {os.path.basename(src)})"
        return ok("convert_document", [row("overwritten" if exists else "converted", out, detail)],
                  notes + list(result_notes or []) + ["The original file was not changed."])

    desc = (f"Convert {real} to {dst_ext} and REPLACE the existing {final}" if exists
            else f"Convert {real} to {final}")
    return guard.submit("convert_document", desc, exists, action, final)